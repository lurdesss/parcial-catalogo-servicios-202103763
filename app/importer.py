"""Importador de CatalogoServicios.xlsx.

Reglas (ver docs/RESOLUCION.md §4):
 R1  Celdas combinadas: el valor de una celda combinada es el de su celda principal, solo dentro de su rango.
 R2  Un servicio N2 se identifica por su código explícito (col. C, ya resuelto por R1). Las filas con el mismo
     código se agrupan en UN servicio; las filas físicas no son servicios.
 R3  Si un atributo del mismo servicio trae valores distintos, gana el primero (orden de filas), se conserva la
     evidencia de todos y se emite una observación. data/mapeo_importacion.json puede fijar un nombre canónico.
 R4  Filas sin código N2 que traen datos de servicio no se asignan a ningún servicio: se omiten y se observan.
 R5  Ausentes se conservan como NULL; nunca se inventa clase, criticidad, tipo, métrica ni ACTIVO.
 R6  Los códigos se conservan como texto; solo se recortan espacios (si cambia, se registra el original).
 R7  Idempotente: upsert por código; no toca asignaciones organizacionales ni la baja lógica.
"""
import hashlib
import json
import re
import unicodedata
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .extensions import db
from .models import Criticality, ImportObservation, ImportRun, ServiceClass, ServiceL1, ServiceL2, ServiceType

SHEET = "Servicios Externos"
HEADER_ROW, DATA_FIRST, DATA_LAST = 4, 5, 101
LIST_FIRST, LIST_LAST = 112, 122
HEADERS = ["COD.N1", "SERVICIO - Nivel 1", "COD.N2", "SERVICIO - Nivel 2", "ACTIVO", "CLASE DE SERVICIO",
           "CRITICIDAD", "TIPO DE SERVICIO", "Descripción", "Métrica", "Minimo", "Maximo"]
(A, B, C, D, E, F, G, H, DESC, J, K, L) = range(1, 13)

CLASSES = ["A DEMANDA", "RECURRENTE"]
CRITICALITIES = ["Very Low", "Low", "Normal", "High", "Very High"]
TYPES = ["Back End", "Demostration", "End User Service", "Front End", "IT Management", "IT Operational",
         "Other", "Project", "Reporting", "Training", "Underpinning Contract"]


class ImportFailure(Exception):
    pass


def norm(s):
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().lower()


def clean(v, keep_number=False):
    if v is None:
        return None
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, (int, float)):
        return v if keep_number else (str(int(v)) if float(v).is_integer() else str(v))
    s = str(v).strip()
    return s or None


def col_letter(c):
    return "ABCDEFGHIJKL"[c - 1]


class Grid:
    """Acceso a celdas resolviendo combinaciones (R1)."""

    def __init__(self, ws):
        self.ws = ws
        self.anchor = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    self.anchor[(r, c)] = (rng.min_row, rng.min_col, str(rng))

    def value(self, r, c):
        a = self.anchor.get((r, c))
        return self.ws.cell(a[0], a[1]).value if a else self.ws.cell(r, c).value

    def ref(self, r, c):
        a = self.anchor.get((r, c))
        return a[2] if a else f"{col_letter(c)}{r}"


class Ctx:
    def __init__(self, run, mapping):
        self.run, self.mapping = run, mapping
        self.obs = []
        self.created = self.updated = self.skipped = 0

    def observe(self, kind, detail, row_ref=None, code=None, evidence=None):
        o = ImportObservation(run_id=self.run.id, kind=kind, sheet=SHEET, row_ref=row_ref, code=code,
                              detail=detail, evidence=json.dumps(evidence, ensure_ascii=False, default=str) if evidence else None)
        db.session.add(o)
        self.obs.append(o)


def _pick(ctx, code, field, entries, canonical=None):
    """R3: devuelve el valor elegido entre [(fila, valor)], observando conflictos."""
    distinct, seen = [], set()
    for row, val in entries:
        key = norm(val) if isinstance(val, str) else repr(val)
        if key not in seen:
            seen.add(key)
            distinct.append((row, val))
    if not distinct:
        return None
    chosen = distinct[0][1]
    if len(distinct) > 1:
        rule = "primer valor en orden de filas"
        if canonical:
            if canonical["nombre"] in [v for _, v in distinct]:
                chosen, rule = canonical["nombre"], "nombre canónico fijado en mapeo_importacion.json: " + canonical["justificacion"]
            else:
                ctx.observe("mapeo_no_aplicable", f"El nombre canónico configurado para {code} no coincide con ningún valor del archivo.",
                            code=code, evidence={"configurado": canonical["nombre"]})
        ctx.observe("conflicto_atributo",
                    f"{field} de {code} tiene {len(distinct)} valores distintos; se aplicó: {rule}.",
                    row_ref=", ".join(f"fila {r}" for r, _ in distinct), code=code,
                    evidence={"valores": [{"fila": r, "valor": v} for r, v in distinct], "elegido": chosen, "regla": rule})
    return chosen


def parse_number(ctx, raw, code, field, row_ref, notes):
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return Decimal(str(raw))
    try:
        d = Decimal(str(raw).strip())
        if d.is_finite():
            return d
    except InvalidOperation:
        pass
    notes.setdefault("valores_no_numericos", {})[field] = raw
    ctx.observe("valor_no_numerico", f"{field} de {code} no es numérico ({raw!r}); se dejó vacío y se conservó el original en las notas.",
                row_ref=row_ref, code=code)
    return None


def _catalog_maps():
    maps = {}
    for key, model in (("class", ServiceClass), ("crit", Criticality), ("type", ServiceType)):
        maps[key] = {norm(c.name): c for c in db.session.scalars(db.select(model))}
    return maps


def ensure_catalogs(ctx, grid):
    """Siembra los catálogos del enunciado y los contrasta con la lista de opciones E112:H122."""
    sheet_lists = {F: [], G: [], H: []}
    for r in range(LIST_FIRST, LIST_LAST + 1):
        for c in sheet_lists:
            v = clean(grid.ws.cell(r, c).value)
            if v:
                sheet_lists[c].append(v)
    ctx.observe("lista_opciones", f"Zona de listas E{LIST_FIRST}:H{LIST_LAST} reconocida; no se importa como servicios.",
                row_ref=f"E{LIST_FIRST}:H{LIST_LAST}",
                evidence={"clase": sheet_lists[F], "criticidad": sheet_lists[G], "tipo": sheet_lists[H]})
    for model, spec, col, label in ((ServiceClass, CLASSES, F, "clase"), (Criticality, CRITICALITIES, G, "criticidad"),
                                    (ServiceType, TYPES, H, "tipo")):
        existing = {norm(c.name) for c in db.session.scalars(db.select(model))}
        for i, name in enumerate(spec):
            if norm(name) not in existing:
                db.session.add(model(name=name, sort_order=i, is_active=True))
                existing.add(norm(name))
        for extra in sheet_lists[col]:
            if norm(extra) not in existing:
                db.session.add(model(name=extra, sort_order=len(existing), is_active=True))
                existing.add(norm(extra))
                ctx.observe("lista_difiere_catalogo", f"El valor «{extra}» de la lista de {label} no estaba en el enunciado; se agregó al catálogo.",
                            row_ref=f"{col_letter(col)}{LIST_FIRST}:{col_letter(col)}{LIST_LAST}", evidence={"valor": extra})
    db.session.flush()


def _label(ctx, maps, key, raw, code, field, row_ref, notes):
    if raw is None:
        return None
    obj = maps[key].get(norm(raw))
    if obj is None:
        ctx.observe("valor_fuera_de_catalogo", f"{field} de {code} = «{raw}» no existe en el catálogo; se dejó sin valor (no se inventa).",
                    row_ref=row_ref, code=code, evidence={"valor": raw})
        notes.setdefault("fuera_de_catalogo", {})[field] = raw
        return None
    if obj.name != raw:
        ctx.observe("mapeo_etiqueta", f"{field} de {code}: «{raw}» → «{obj.name}».", row_ref=row_ref, code=code,
                    evidence={"original": raw, "normalizado": obj.name})
        notes.setdefault("mapeo_etiquetas", {})[field] = {"original": raw, "normalizado": obj.name}
    return obj.id


def _code(ctx, raw_value, row_ref, level):
    """R6: texto; solo se recortan espacios."""
    if raw_value is None:
        return None, None
    original = raw_value if isinstance(raw_value, str) else str(raw_value)
    code = clean(raw_value)
    if code is None:
        return None, None
    if not isinstance(raw_value, str):
        ctx.observe("codigo_no_texto", f"Código {level} numérico {raw_value!r} convertido a texto «{code}».", row_ref=row_ref, code=code)
    elif code != original:
        ctx.observe("mapeo_codigo", f"Código {level} «{original}» normalizado a «{code}» (espacios).", row_ref=row_ref, code=code,
                    evidence={"original": original, "normalizado": code})
    return code, original


def _upsert(ctx, model, code, values):
    obj = db.session.scalar(db.select(model).where(model.code == code))
    if obj is None:
        db.session.add(model(code=code, is_active=True, **values))
        ctx.created += 1
        return
    changed = {k: v for k, v in values.items() if getattr(obj, k) != v}
    if changed:
        for k, v in changed.items():
            setattr(obj, k, v)
        ctx.updated += 1
    else:
        ctx.skipped += 1


def _rng(ranges):
    return ", ".join(sorted(ranges, key=lambda x: int(re.search(r"\d+", x).group())))


def _span(rows):
    return f"filas {rows[0]}-{rows[-1]}" if len(rows) > 1 and rows[-1] - rows[0] == len(rows) - 1 else \
        "filas " + ", ".join(map(str, rows))


def read_groups(ctx, grid):
    l1, l2 = {}, {}
    for r in range(DATA_FIRST, DATA_LAST + 1):
        cells = {c: grid.value(r, c) for c in range(1, 13)}
        if all(clean(v) is None for v in cells.values()):
            continue
        a_code, a_orig = _code(ctx, cells[A], f"A{r}", "N1")
        c_code, c_orig = _code(ctx, cells[C], f"C{r}", "N2")
        if a_code is None and c_code is not None and "." in c_code:   # R2 (prefijo explícito del código N2)
            prefix = c_code.rsplit(".", 1)[0]
            if prefix in l1:
                a_code, a_orig = prefix, prefix
                ctx.observe("n1_inferido_por_prefijo", f"Fila {r}: sin COD.N1; se usó el prefijo «{prefix}» del código N2 (ya existente).",
                            row_ref=f"fila {r}", code=c_code)
        if a_code is not None:
            g = l1.setdefault(a_code, {"orig": a_orig, "rows": [], "names": [], "ranges": set()})
            g["rows"].append(r)
            g["ranges"].add(grid.ref(r, A))
            if clean(cells[B]):
                g["names"].append((r, clean(cells[B])))
        if c_code is None:
            payload = [col_letter(c) for c in range(D, L + 1) if clean(cells[c]) is not None]
            if payload:   # R4
                ctx.observe("fila_sin_codigo",
                            f"Fila {r} sin COD.N2 y fuera de una combinación; trae datos en {', '.join(payload)}. No se asignó a ningún servicio.",
                            row_ref=f"fila {r}", evidence={col_letter(c): clean(cells[c], True) for c in range(D, L + 1) if clean(cells[c], True) is not None})
            continue
        g = l2.setdefault(c_code, {"orig": c_orig, "rows": [], "l1": [], "f": {c: [] for c in range(D, L + 1)}, "ranges": set()})
        g["rows"].append(r)
        g["ranges"].add(grid.ref(r, C))
        if a_code:
            g["l1"].append((r, a_code))
        for c in range(D, L + 1):
            v = clean(cells[c], keep_number=c in (K, L))
            if v is not None:
                g["f"][c].append((r, v))
    for r in range(DATA_LAST + 1, LIST_FIRST - 1):   # entre los datos y el título de las listas (fila 111)
        payload = [col_letter(c) for c in range(1, 13) if clean(grid.value(r, c)) is not None]
        if payload:
            ctx.observe("fila_fuera_de_rango", f"Fila {r} trae datos en {', '.join(payload)} fuera del rango de datos "
                        f"{DATA_FIRST}-{DATA_LAST}; no se importó.", row_ref=f"fila {r}")
    return l1, l2


def run_import(xlsx_path, mapping_path=None):
    path = Path(xlsx_path)
    if not path.is_file():
        raise ImportFailure(f"No se encontró el archivo {path}. Colóquelo en data/CatalogoServicios.xlsx.")
    mapping = {"nombres_canonicos": {}}
    if mapping_path and Path(mapping_path).is_file():
        mapping = json.loads(Path(mapping_path).read_text(encoding="utf-8"))
    digest = hashlib.sha256(path.read_bytes()).hexdigest()

    run = ImportRun(file_name=path.name, file_sha256=digest, status="running")
    db.session.add(run)
    db.session.flush()
    ctx = Ctx(run, mapping)
    try:
        wb = load_workbook(path, data_only=True)   # solo lectura; el archivo original no se modifica
        if SHEET not in wb.sheetnames:
            raise ImportFailure(f"La hoja «{SHEET}» no existe (hojas: {wb.sheetnames}).")
        ws = wb[SHEET]
        found = [ws.cell(HEADER_ROW, c).value for c in range(1, 13)]
        if [norm(x or "") for x in found] != [norm(h) for h in HEADERS]:
            raise ImportFailure(f"Los encabezados de A{HEADER_ROW}:L{HEADER_ROW} no coinciden con lo esperado: {found}")
        grid = Grid(ws)
        ensure_catalogs(ctx, grid)
        maps = _catalog_maps()
        l1, l2 = read_groups(ctx, grid)
        canon = mapping.get("nombres_canonicos", {})

        for code, g in l1.items():
            name = _pick(ctx, code, "Nombre N1", g["names"], canon.get(code))
            notes = {"nombres_encontrados": [{"fila": r, "valor": v} for r, v in g["names"]]} if len({norm(v) for _, v in g["names"]}) > 1 else {}
            if name is None:
                ctx.observe("nombre_ausente", f"El código N1 {code} no tiene nombre; se usó el código como nombre provisional.", code=code)
                name, notes["nombre_provisional"] = code, True
            _upsert(ctx, ServiceL1, code, dict(
                code_original=g["orig"], name=name, source_sheet=SHEET,
                source_range=_rng(g["ranges"]) + f" ({_span(g['rows'])})",
                source_notes=json.dumps(notes, ensure_ascii=False) if notes else None))
        db.session.flush()
        l1_ids = {s.code: s.id for s in db.session.scalars(db.select(ServiceL1))}

        for code, g in l2.items():
            notes, reasons = {}, []
            row_ref = _span(g["rows"])
            parents = _pick(ctx, code, "Servicio N1 padre", g["l1"])
            if parents is None:
                ctx.observe("n2_sin_n1", f"El servicio {code} no tiene COD.N1 resoluble; no se importó.", row_ref=row_ref, code=code)
                continue
            name = _pick(ctx, code, "Nombre N2", g["f"][D], canon.get(code))
            if name is None:
                ctx.observe("nombre_ausente", f"El código N2 {code} no tiene nombre; se usó el código como nombre provisional.", row_ref=row_ref, code=code)
                name, reasons = code, reasons + ["nombre provisional"]
            raw = {c: _pick(ctx, code, HEADERS[c - 1], g["f"][c]) for c in range(E, L + 1)}
            activo_raw = raw[E]
            activo = None
            if activo_raw is not None and str(activo_raw).strip().upper() in ("S", "N"):
                activo, activo_raw = str(activo_raw).strip().upper(), None
            elif activo_raw is not None:
                ctx.observe("activo_desconocido", f"ACTIVO de {code} = «{activo_raw}» no es S/N; se conserva como desconocido.",
                            row_ref=row_ref, code=code, evidence={"valor": activo_raw})
            ids = {
                "class_id": _label(ctx, maps, "class", raw[F], code, "CLASE DE SERVICIO", row_ref, notes),
                "criticality_id": _label(ctx, maps, "crit", raw[G], code, "CRITICIDAD", row_ref, notes),
                "type_id": _label(ctx, maps, "type", raw[H], code, "TIPO DE SERVICIO", row_ref, notes),
            }
            vmin = parse_number(ctx, raw[K], code, "Minimo", row_ref, notes)
            vmax = parse_number(ctx, raw[L], code, "Maximo", row_ref, notes)
            if vmin is not None and vmax is not None and vmin > vmax:
                ctx.observe("umbral_invalido", f"{code}: Minimo ({vmin}) > Maximo ({vmax}); se conservan vacíos y los originales en notas.",
                            row_ref=row_ref, code=code)
                notes["umbral_original"] = {"Minimo": str(vmin), "Maximo": str(vmax)}
                vmin = vmax = None
            missing = [lbl for lbl, val in (("ACTIVO", activo), ("CLASE", ids["class_id"]), ("CRITICIDAD", ids["criticality_id"]),
                                            ("TIPO", ids["type_id"])) if val is None]
            if missing:
                reasons.append("sin " + ", ".join(missing))
                ctx.observe("atributos_incompletos", f"{code}: faltan {', '.join(missing)}; se importó con valores desconocidos y estado «revisar».",
                            row_ref=row_ref, code=code)
            if any(o.kind == "conflicto_atributo" and o.code == code for o in ctx.obs):
                reasons.append("conflicto de atributos")
            if reasons:
                notes["motivos_revision"] = reasons
            if len(g["rows"]) > 1:
                notes["filas_agrupadas"] = g["rows"]
            _upsert(ctx, ServiceL2, code, dict(
                code_original=g["orig"], name=name, level1_id=l1_ids[parents], activo_excel=activo,
                activo_raw=str(activo_raw) if activo_raw is not None else None, description=raw[DESC], metric=raw[J],
                min_value=vmin, max_value=vmax, review_status="revisar" if reasons else "ok",
                source_sheet=SHEET, source_range=_rng(g["ranges"]) + f" ({row_ref})",
                source_notes=json.dumps(notes, ensure_ascii=False) if notes else None, **ids))

        run.created, run.updated, run.skipped = ctx.created, ctx.updated, ctx.skipped
        run.observed = len(ctx.obs)
        run.n1_codes, run.n2_codes = len(l1), len(l2)
        by_kind = {}
        for o in ctx.obs:
            by_kind[o.kind] = by_kind.get(o.kind, 0) + 1
        run.summary = json.dumps({"observaciones_por_tipo": by_kind,
                                  "control": {"n1_esperados": 12, "n2_esperados": 46,
                                              "n1_leidos": len(l1), "n2_leidos": len(l2),
                                              "coincide": len(l1) == 12 and len(l2) == 46}}, ensure_ascii=False)
        run.status = "ok"
        run.finished_at = datetime.now(timezone.utc).replace(tzinfo=None)
        db.session.commit()
        return run
    except Exception as e:
        db.session.rollback()
        failed = ImportRun(file_name=path.name, file_sha256=digest, status="error", summary=json.dumps({"error": str(e)}),
                           finished_at=datetime.now(timezone.utc).replace(tzinfo=None))
        db.session.add(failed)
        db.session.commit()
        if isinstance(e, ImportFailure):
            raise
        raise ImportFailure(f"Error inesperado durante la importación: {e}") from e
