"""Libro SINTÉTICO que imita la estructura descrita en el enunciado (NO es el archivo real).

Sirve para probar las reglas del importador de forma controlada: celdas combinadas, conflicto SE.12,
códigos como texto, atributos ausentes, filas de continuación sin código y listas E112:H122.
"""
from openpyxl import Workbook

HEADERS = ["COD.N1", "SERVICIO - Nivel 1", "COD.N2", "SERVICIO - Nivel 2", "ACTIVO", "CLASE DE SERVICIO",
           "CRITICIDAD", "TIPO DE SERVICIO", "Descripción", "Métrica", "Minimo", "Maximo"]


def build(path, corrupt_headers=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Servicios Externos"
    for i, h in enumerate(HEADERS, start=1):
        ws.cell(4, i, "OTRO" if corrupt_headers and i == 1 else h)

    def row(r, *vals):
        for i, v in enumerate(vals, start=1):
            if v is not None:
                ws.cell(r, i, v)

    # SE.01: N1 combinado en filas 5-8; SE.01.1 combinado en 5-6 (atributos combinados); fila 8 sin código
    row(5, "SE.01", "Gestionar Infraestructura", "SE.01.1", "Monitorear Servidores", "S", "RECURRENTE", "High",
        "Back End", "Monitoreo continuo", "Disponibilidad", 95, 99.9)
    row(6, None, None, None, None, None, None, None, None, None, "Disponibilidad distinta", None, None)  # conflicto de métrica
    ws.merge_cells("A5:A8"); ws.merge_cells("B5:B8")
    ws.merge_cells("C5:C6"); ws.merge_cells("D5:D6")
    ws.merge_cells("E5:E6"); ws.merge_cells("F5:F6"); ws.merge_cells("G5:G6"); ws.merge_cells("H5:H6")
    row(7, None, None, "SE.01.2", "Respaldar Datos", "N", "a demanda", "Normal", "IT Operational", None, None, None, None)
    row(8, None, None, None, "Fila huérfana de continuación", None, None, None, None, "texto suelto", None, None, None)
    # SE.02: umbral inválido (min > max) y valor fuera de catálogo
    row(9, "SE.02", "Soporte a Usuarios", "SE.02.1", "Atender Incidentes", "S", "RECURRENTE", "Muy Alta", "Front End",
        None, "Tiempo", 10, 5)
    row(10, "SE.02", "Soporte a Usuarios", " SE.02.2", "Atender Requerimientos", "S", "A DEMANDA", "Low", "Other",
        None, None, None, None)
    # SE.12: conflicto de nombre N1 en filas 99/100 y atributos ausentes en 99-101
    row(99, "SE.12", "Suministrar Analitica", "SE.12.1", "Generar Reportes", None, None, None, None, None, None, None, None)
    row(100, "SE.12", "Mantener Tableros de Control", "SE.12.2", "Mantener Tableros", "?", None, None, None, None, None, None, None)
    row(101, "SE.12", "Suministrar Analitica", "SE.12.3", "Publicar Indicadores", None, None, None, None, None, None, None, None)
    # Lista de opciones E112:H122
    for i, v in enumerate(["S", "N"]):
        ws.cell(112 + i, 5, v)
    for i, v in enumerate(["A DEMANDA", "RECURRENTE"]):
        ws.cell(112 + i, 6, v)
    for i, v in enumerate(["Very Low", "Low", "Normal", "High", "Very High"]):
        ws.cell(112 + i, 7, v)
    for i, v in enumerate(["Back End", "Demostration", "End User Service", "Front End", "IT Management",
                           "IT Operational", "Other", "Project", "Reporting", "Training", "Underpinning Contract"]):
        ws.cell(112 + i, 8, v)
    wb.save(path)
    return path
